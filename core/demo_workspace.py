"""A self-contained, deterministic demonstration with no customer data or API.

The labels below are authored fixtures, not model predictions. This module only
reads its packaged synthetic dataset; it never opens history stores, originals,
credentials, or network clients. Every returned report is a fresh object.
"""
from __future__ import annotations

from collections import Counter
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
from xml.etree import ElementTree
import zipfile

import pandas as pd

from core.analysis_agent import AnalysisAgentResult, AgentPoint, report_fingerprint
from core.customer_report import build_customer_report
from core.history_workflow import attach_historical_context, date_scope
from models.models import FeedbackResult


_FIXTURE = Path(__file__).resolve().parent / "demo_data" / "reviews.json"
_SOURCE_COLUMNS = ["评价编号", "商品", "评价日期", "顾客评价"]
_DEMO_SCOPE = "离线示例回放：数据为固定合成评价，标签和回答由程序预置或计算，未调用 LLM；仅展示这两批示例，不能推断真实业务表现或因果效果。"


def _load_fixture() -> dict:
    payload = json.loads(_FIXTURE.read_text(encoding="utf-8"))
    if payload.get("data_origin") != "synthetic" or payload.get("version") != "demo_workspace_v1":
        raise ValueError("示例数据来源或版本不匹配。")
    return payload


def _source_frame(batch: dict) -> pd.DataFrame:
    frame = pd.DataFrame([
        {"评价编号": item["review_id"], "商品": item["product"],
         "评价日期": item["date"], "顾客评价": item["text"]}
        for item in batch["reviews"]
    ], columns=_SOURCE_COLUMNS)
    frame.attrs["source_rows"] = list(range(2, len(frame) + 2))
    return frame


def _build_report(batch: dict, metadata: dict) -> dict:
    source = _source_frame(batch)
    results = []
    for number, item in enumerate(batch["reviews"], 2):
        result = FeedbackResult(
            feedback_id=item["review_id"], source_row=number, original_text=item["text"],
            sentiment=item["sentiment"], category=item["category"], priority=item["priority"],
            summary=item["summary"], processing_status="success", latency_ms=0,
            aspects=item["aspects"],
        )
        results.append(result.model_dump(mode="json"))
    report = build_customer_report(
        source, pd.DataFrame(results), "顾客评价", "商品",
        analysis_meta={"prompt_version": "demo_predefined_v1", "mode": "offline_fixture",
                       "model_calls": 0, "actual_requests": 0,
                       "label_origin": "predefined", "data_origin": "synthetic",
                       "demo_version": metadata["version"]},
    )
    # The old report contract calls successful classifications "model success".
    # Keep its invariant for existing consumers but make all user-facing sources
    # explicit: these classifications were loaded, never sent to a model.
    names = report["detail_columns"]
    report["detail"][names["result_source"]] = "离线预置标签（合成示例，未调用模型）"
    report["negative_examples"][names["result_source"]] = "离线预置标签（合成示例，未调用模型）"
    report["notes"] = [
        "本报告为合成示例；评价、分类、优先级和主题由开发者预置，统计由程序计算；本次未产生模型请求。",
        f"原始记录 {len(source)} 条，全部带有可复核的预置标签。模型调用数为 0，不能用示例推断模型质量。",
        *[note for note in report["notes"] if "模型" not in note and not note.startswith("原始记录")],
    ]
    report.update(
        period=date_scope(source, "评价日期"),
        data_origin="synthetic", answer_origin="fallback",
        provenance={"data_origin": "synthetic", "label_origin": "predefined",
                    "answer_origin": "fallback", "model_calls": 0,
                    "dataset_version": metadata["version"], "batch_id": batch["id"],
                    "filename": batch["filename"], "scope_note": _DEMO_SCOPE},
    )
    return report


def _replay(report: dict, previous: dict) -> tuple[list[str], list[dict]]:
    """Compute two replay turns from the same traceable report used by charts."""
    issues = [topic for topic in report["topics"] if topic["kind"] == "issue"]
    top = max(issues, key=lambda item: (item["review_count"], item["topic"]))
    question_one = "这批评价最值得优先处理什么？"
    question_two = f"{top['topic']}集中在哪个商品？和上批相比有什么变化？"
    names = report["detail_columns"]
    detail = report["detail"].set_index(names["source_row"])
    counts = Counter(str(detail.loc[number, report["group_column"]]) for number in top["source_rows"])
    product, product_count = sorted(counts.items(), key=lambda item: (-item[1], item[0]))[0]
    current_rows = [number for number in top["source_rows"] if str(detail.loc[number, report["group_column"]]) == product]
    comparison = report["historical_context"]["comparison"]
    topic_change = next(item for item in comparison["topics"] if item["key"] == top["key"])
    old_top = next((item for item in previous["report"]["topics"] if item["key"] == top["key"]), None)
    history_rows = old_top["source_rows"][:3] if old_top else []
    old_names = previous["report"]["detail_columns"]
    old_detail = previous["report"]["detail"].set_index(old_names["source_row"])
    history_evidence = [
        {"batch_id": previous["id"], "filename": previous["filename"], "source_row": number,
         "text": str(old_detail.loc[number, previous["report"]["text_column"]])}
        for number in history_rows
    ]
    base = {
        "status": "completed", "answer_origin": "fallback", "answer_type": "supported",
        "scope_note": _DEMO_SCOPE, "model_calls": 0, "tool_calls": 0,
        "tool_trace": [], "total_tokens": None, "usage_reported_requests": 0,
        "fallback_reason": "离线预置示例回放，未调用模型。",
        "report_fingerprint": report_fingerprint(report),
    }
    top_stat = f"本批有效评价 {report['metrics']['success_rows']} 条；{top['topic']} {top['review_count']} 条，占 {top['share']:.1%}。"
    first = AnalysisAgentResult(
        **base, question=question_one,
        conclusion=f"先核对「{top['topic']}」。它是本批出现最多的问题主题，共 {top['review_count']} 条评价提及（占有效评价 {top['share']:.1%}）。这是排查线索，不代表已经确认故障原因。",
        conclusion_metrics=[top_stat],
        findings=[AgentPoint(title=f"优先排查：{top['topic']}",
                             narrative="先看相关评价原文，再核对订单与履约信息；主题提及次数不等于已核实事故数。",
                             metrics=[top_stat], evidence_rows=top["source_rows"][:3])],
        actions=[AgentPoint(title="按商品查原文，再安排责任人核查",
                            narrative="下一轮可追问问题集中在哪个商品，以及与上批的差异。处理后继续收集同口径评价，不能只凭两批评价判断措施有效。",
                            metrics=[], evidence_rows=top["source_rows"][:3])],
        evidence_rows=top["source_rows"][:3],
    )
    change = topic_change["change_pp"]
    direction = "上升" if change > 0 else "下降" if change < 0 else "持平"
    delta_phrase = f"{direction} {abs(change):.1f} 个百分点" if change else "持平"
    batch_stat = (f"{top['topic']}：本批 {topic_change['current_count']}/{report['metrics']['success_rows']} 条（{topic_change['current_share']:.1%}），"
                  f"上批 {topic_change['previous_count']}/{previous['report']['metrics']['success_rows']} 条（{topic_change['previous_share']:.1%}），占比{delta_phrase}。")
    product_stat = f"本批 {top['review_count']} 条「{top['topic']}」评价中，{product}有 {product_count} 条。"
    current_references = "、".join(str(number) for number in current_rows[:3])
    old_references = "、".join(str(number) for number in history_rows)
    second = AnalysisAgentResult(
        **base, question=question_two,
        conclusion=(f"本批「{top['topic']}」最集中在 {product}，占该问题评价 {product_count}/{top['review_count']} 条。"
                    f"与上批相比，该主题占有效评价的比例{delta_phrase}。可核对本批原表第 {current_references} 行"
                    f"与上批《{previous['filename']}》第 {old_references} 行；两个批次的行号独立。"
                    "建议先核对该商品的发货节点、物流记录与承诺时效，再决定改进动作；示例变化不能证明因果。"),
        conclusion_metrics=[product_stat, batch_stat],
        findings=[AgentPoint(title="商品集中与批次变化", narrative="结合商品维度和两批同口径原文核查，而不是把全部物流评价平均分配给各商品。",
                             metrics=[product_stat, batch_stat], evidence_rows=current_rows[:3])],
        actions=[AgentPoint(title="核对履约环节", narrative="由运营按原文关联订单，检查出库等待与运输等待；对异常订单给出进度提醒。此建议是待验证的排查步骤。",
                            metrics=[], evidence_rows=current_rows[:3])],
        evidence_rows=current_rows[:3], historical_evidence=history_evidence,
    )
    return [question_one, question_two], [first.model_dump(mode="json"), second.model_dump(mode="json")]


def _workbook_bytes(payload: dict) -> bytes:
    """Generate an uploadable source workbook using the app's Excel dependency."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "CustomerInsight 合成评价示例"
    workbook.properties.description = "两批各 30 条固定合成评价；非真实顾客数据；不含模型预测或人工金标准。"
    workbook.properties.created = datetime(2026, 10, 15, tzinfo=timezone.utc)
    workbook.properties.modified = datetime(2026, 10, 15, tzinfo=timezone.utc)
    for key, title in (("current", "本批评价"), ("previous", "上批评价")):
        sheet = workbook.create_sheet(title)
        sheet.append(_SOURCE_COLUMNS)
        frame = _source_frame(payload[key])
        for row in frame.itertuples(index=False, name=None):
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 24
        for cell in sheet[1]:
            cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="287D88")
        for column, width in (("A", 21), ("B", 24), ("C", 16), ("D", 85)):
            sheet.column_dimensions[column].width = width
        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height = 36
            for cell in row:
                cell.font = Font(name="Microsoft YaHei", size=11)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F0F7F8")
    note = workbook.create_sheet("来源说明")
    note.append(["项目", "说明"])
    for row in (("数据来源", "固定合成示例，不是真实顾客数据"),
                ("数据版本", payload["version"]),
                ("行号约定", "每张评价表第一行为表头，数据原始行号从 2 开始；两批独立编号"),
                ("标签来源", "界面内置分类和主题由开发者预置，不是模型预测或人工金标准；原表不含标签列"),
                ("模型调用", "打开示例和查看预置回放均为 0 次；自由追问需明确使用实时分析服务"),
                ("使用限制", "仅用于演示流程、图表联动和原文查证；不能作为真实业务结论或模型质量评测")):
        note.append(row)
    note.column_dimensions["A"].width = 18
    note.column_dimensions["B"].width = 100
    for row in note.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        note.row_dimensions[row[0].row].height = 32
    stream = io.BytesIO()
    workbook.save(stream)
    # Stable ZIP/core timestamps preserve content identity across repeated demo
    # imports. Otherwise the history store sees a new file hash every run.
    stable = io.BytesIO()
    with zipfile.ZipFile(stream) as source, zipfile.ZipFile(stable, "w") as target:
        for member in source.infolist():
            content = source.read(member.filename)
            if member.filename == "docProps/core.xml":
                root = ElementTree.fromstring(content)
                for field in ("created", "modified"):
                    node = root.find("{http://purl.org/dc/terms/}" + field)
                    if node is not None:
                        node.text = "2026-10-15T00:00:00Z"
                content = ElementTree.tostring(root, encoding="utf-8")
            member.date_time = (2026, 10, 15, 0, 0, 0)
            target.writestr(member, content)
    return stable.getvalue()


def load_demo_workspace() -> dict:
    """Return a complete fresh workspace; no history or environment is consulted.

    ``current_report`` already contains the selected ``previous_batch`` context.
    ``replay_turns`` conform to AnalysisAgentResult and are explicitly fallback
    answers, not model output. ``workbook_bytes`` contains raw source records only.
    """
    payload = _load_fixture()
    previous_report = _build_report(payload["previous"], payload)
    current_report = _build_report(payload["current"], payload)
    previous_batch = {key: payload["previous"][key] for key in ("id", "filename", "created_at")}
    previous_batch["report"] = previous_report
    current_report = attach_historical_context(current_report, previous_batch)
    questions, turns = _replay(current_report, previous_batch)
    metadata = {
        "version": payload["version"], "data_origin": "synthetic", "label_origin": "predefined",
        "demo_version": payload["version"], "actual_requests": 0,
        "answer_origin": "fallback", "model_calls": 0, "scope_note": _DEMO_SCOPE,
        "title": "青岚杯具 · 两批评价演示", "current_rows": len(payload["current"]["reviews"]),
        "previous_rows": len(payload["previous"]["reviews"]),
        "current_filename": payload["current"]["filename"],
        "previous_filename": payload["previous"]["filename"],
        "description": "查看本批问题，点击图表核对原文，再体验两轮离线回放或明确发起实时追问。",
    }
    return {
        "data_origin": "synthetic", "answer_origin": "fallback",
        "current_report": current_report, "previous_report": previous_report,
        "previous_batch": previous_batch, "questions": questions, "replay_turns": turns,
        "current_mapping": {"sheet_name": "本批评价", "text_column": "顾客评价",
                            "group_column": "商品", "date_column": "评价日期", "encoding": None},
        "previous_mapping": {"sheet_name": "上批评价", "text_column": "顾客评价",
                             "group_column": "商品", "date_column": "评价日期", "encoding": None},
        "workbook_bytes": _workbook_bytes(payload),
        "workbook_filename": "CustomerInsight_合成评价示例_两批60条.xlsx", "metadata": metadata,
    }


def save_demo_workspace(history) -> dict:
    """Explicitly import fixtures into the caller's already scoped history.

    No global store, account, database configuration or client is resolved here.
    Existing matching imports are reused without rewriting reports or chats.
    Persisted reports contain no fixture batch links; callers explicitly bind
    the returned previous ID when asking questions or comparing saved batches.
    """
    workspace = load_demo_workspace()
    file_bytes = workspace["workbook_bytes"]
    digest = hashlib.sha256(file_bytes).hexdigest()
    identifiers = {}
    for kind in ("previous", "current"):
        mapping = workspace[f"{kind}_mapping"]
        existing = history.find_batch(digest, mapping)
        if existing is not None:
            identifier = existing["id"]
        else:
            report = workspace[f"{kind}_report"]
            report.pop("historical_context", None)
            identifier = history.save_batch(
                report, filename=workspace["metadata"][f"{kind}_filename"],
                file_bytes=file_bytes, mapping=mapping,
            )
        identifiers[f"{kind}_batch_id"] = identifier
    return identifiers
