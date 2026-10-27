"""Bounded .xlsx ingestion with physical worksheet row identities.

The first worksheet row is the header. Blank rows between records are retained,
so the existing feedback preparation contract still refers to real Excel rows.
"""

import io
import re
from pathlib import PurePath
from zipfile import BadZipFile, ZipFile
from xml.etree import ElementTree

import pandas as pd

from core.config import Config
from core.csv_io import CsvInputError, _row_limit, _validate_columns, prepare_feedback_rows
from models.enums import ProcessingStatus
from models.models import FeedbackResult

MAX_WORKSHEETS = 20
MAX_COLUMNS = 100
MAX_UNCOMPRESSED_BYTES = 40 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 1000
_CELL_REFERENCE = re.compile(r"^([A-Z]+)([1-9][0-9]*)$")
_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


class ExcelInputError(CsvInputError):
    """An uploaded workbook cannot be safely processed."""


def _column_number(letters: str) -> int:
    value = 0
    for character in letters:
        value = value * 26 + ord(character) - ord("A") + 1
    return value


def _coordinate(value: str) -> tuple[int, int]:
    match = _CELL_REFERENCE.fullmatch(value)
    if not match:
        raise ExcelInputError("Excel 单元格坐标无效，请重新保存为 .xlsx。")
    return int(match[2]), _column_number(match[1])


def _read_xml(archive: ZipFile, name: str) -> ElementTree.Element:
    content = archive.read(name)
    # XML also permits UTF-16/32; strip their ASCII interleaving NULs so a DTD
    # cannot evade the pre-parser check by changing the document encoding.
    upper = content.replace(b"\x00", b"").upper()
    if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
        raise ExcelInputError("Excel 包含不支持的 XML 声明，请重新导出为普通 .xlsx。")
    return ElementTree.fromstring(content)


def _inspect_archive(data: bytes, filename: str | None, *, max_rows: int | None = None) -> list[str]:
    limit = _row_limit(max_rows)
    if not isinstance(data, bytes) or not data:
        raise ExcelInputError("请上传非空的 Excel .xlsx 文件。")
    if len(data) > Config.MAX_CSV_BYTES:
        raise ExcelInputError("Excel 文件超过 5 MB 上限，请拆分文件。")
    if filename is not None:
        if not isinstance(filename, str):
            raise ExcelInputError("Excel 文件名无效。")
        suffix = PurePath(filename).suffix.lower()
        if suffix == ".xls":
            raise ExcelInputError("暂不支持旧版 .xls，请在 Excel 中另存为 .xlsx 后上传。")
        if suffix != ".xlsx":
            raise ExcelInputError("请上传扩展名为 .xlsx 的 Excel 文件。")
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            names = [entry.filename for entry in entries]
            if len(entries) > MAX_ARCHIVE_ENTRIES or len(set(names)) != len(names):
                raise ExcelInputError("Excel 内部文件数量异常或包含重复项目，请重新保存。")
            if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
                raise ExcelInputError("Excel 解压后内容超过 40 MB 安全上限，请简化或拆分文件。")
            if any(entry.flag_bits & 1 for entry in entries):
                raise ExcelInputError("不支持加密的 Excel 文件，请去除密码后重新上传。")
            if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
                raise ExcelInputError("文件不是有效的 Excel .xlsx 工作簿。")
            if any("vbaproject" in name.lower() or "externallinks/" in name.lower() for name in names):
                raise ExcelInputError("请移除宏和外部工作簿链接，再上传普通 .xlsx 文件。")
            workbook = _read_xml(archive, "xl/workbook.xml")
            sheets = workbook.findall(f"{{{_MAIN_NS}}}sheets/{{{_MAIN_NS}}}sheet")
            sheet_names = [sheet.get("name", "") for sheet in sheets]
            if not sheet_names or len(sheet_names) > MAX_WORKSHEETS:
                raise ExcelInputError(f"工作簿必须包含 1 至 {MAX_WORKSHEETS} 个工作表，请拆分文件。")
            if any(not name.strip() for name in sheet_names) or len(set(sheet_names)) != len(sheet_names):
                raise ExcelInputError("Excel 工作表名称为空或重复，请重新保存。")
            # Scan XML before the Excel library reads it. Both the declared and
            # actual coordinates are checked; forged dimensions cannot bypass limits.
            for name in names:
                if not name.lower().endswith(".xml"):
                    continue
                root = _read_xml(archive, name)
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    declared_bounds = None
                    actual_row, actual_column, row_elements = 1, 1, 0
                    seen_cells = set()
                    for element in root.iter():
                        if element.tag == f"{{{_MAIN_NS}}}dimension":
                            for reference in element.get("ref", "A1").split(":"):
                                row, column = _coordinate(reference)
                                if row > limit + 1 or column > MAX_COLUMNS:
                                    raise ExcelInputError(f"每张工作表最多 {limit} 条数据、{MAX_COLUMNS} 列；请删除多余格式或拆分文件。")
                            declared_bounds = _coordinate(element.get("ref", "A1").split(":")[-1])
                        elif element.tag == f"{{{_MAIN_NS}}}c":
                            reference = element.get("r", "")
                            row, column = _coordinate(reference)
                            if row > limit + 1 or column > MAX_COLUMNS:
                                raise ExcelInputError(f"每张工作表最多 {limit} 条数据、{MAX_COLUMNS} 列，请拆分文件。")
                            if reference in seen_cells:
                                raise ExcelInputError("Excel 包含重复单元格坐标，请重新保存文件。")
                            seen_cells.add(reference)
                            actual_row, actual_column = max(actual_row, row), max(actual_column, column)
                        elif element.tag == f"{{{_MAIN_NS}}}row":
                            try:
                                row = int(element.get("r", "0"))
                            except ValueError:
                                raise ExcelInputError("Excel 行号无效，请重新保存文件。") from None
                            if row < 1 or row > limit + 1:
                                raise ExcelInputError(f"每张工作表最多 {limit} 条数据，请拆分文件。")
                            row_elements += 1
                            if row_elements > limit + 1:
                                raise ExcelInputError("Excel 行记录数量异常，请重新保存文件。")
                            actual_row = max(actual_row, row)
                    if declared_bounds and (actual_row > declared_bounds[0] or actual_column > declared_bounds[1]):
                        raise ExcelInputError("Excel 工作表声明范围与实际数据不一致，请重新保存文件。")
            return sheet_names
    except ExcelInputError:
        raise
    except (BadZipFile, KeyError, ElementTree.ParseError, OSError, RuntimeError, ValueError) as exc:
        raise ExcelInputError("Excel 文件已损坏或格式无效，请重新导出为 .xlsx。") from exc


def inspect_workbook(data: bytes, filename: str | None = None, *, max_rows: int | None = None) -> list[str]:
    """Return worksheet names after bounded ZIP/XML inspection, without AI calls."""
    return _inspect_archive(data, filename, max_rows=max_rows)


def read_feedback_xlsx(
    data: bytes, filename: str | None = None, sheet_name: str | None = None,
    *, max_rows: int | None = None,
) -> pd.DataFrame:
    """Read a selected worksheet, preserving raw cells and formula provenance."""
    limit = _row_limit(max_rows)
    sheet_names = _inspect_archive(data, filename, max_rows=limit)
    selected = sheet_names[0] if sheet_name is None else sheet_name
    if selected not in sheet_names:
        raise ExcelInputError("所选工作表不存在，请重新选择。")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ExcelInputError("Excel 读取依赖尚未安装，请安装项目依赖后重试。") from exc
    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        sheet = workbook[selected]
        if sheet.max_row is None or sheet.max_column is None:
            sheet.calculate_dimension(force=True)
        if sheet.max_row > limit + 1 or sheet.max_column > MAX_COLUMNS:
            raise ExcelInputError(f"所选工作表最多 {limit} 条数据、{MAX_COLUMNS} 列。")
        cells = list(sheet.iter_rows())
        if not cells:
            raise ExcelInputError("工作表为空，请在第一行填写表头。")
        header = [cell.value for cell in cells[0]]
        try:
            _validate_columns(header)
        except CsvInputError as exc:
            raise ExcelInputError(str(exc).replace("CSV", "Excel")) from exc
        if any(cell.data_type == "f" for cell in cells[0]):
            raise ExcelInputError("Excel 表头不能使用公式，请粘贴为文本。")
        if len(cells) < 2:
            raise ExcelInputError("工作表只有表头，没有顾客评价记录。")
        values, formula_cells, error_cells = [], [], []
        for row_number, row in enumerate(cells[1:], start=2):
            values.append([cell.value for cell in row])
            for column, cell in zip(header, row):
                if cell.data_type == "f":
                    formula_cells.append((row_number, column))
                elif cell.data_type == "e":
                    error_cells.append((row_number, column))
        frame = pd.DataFrame(values, columns=header, dtype=object)
        frame.attrs.update(
            source_rows=list(range(2, len(frame) + 2)),
            formula_cells=formula_cells,
            error_cells=error_cells,
            sheet_name=selected,
            input_format="xlsx",
        )
        return frame
    except ExcelInputError:
        raise
    except Exception as exc:
        raise ExcelInputError("Excel 无法读取，请取消合并表头、移除损坏内容后重新保存。") from exc
    finally:
        if workbook is not None:
            workbook.close()


def prepare_excel_feedback_rows(frame: pd.DataFrame, text_column: str, *, max_rows: int | None = None) -> list[FeedbackResult]:
    """Reuse text validation/deduplication, while rejecting formula/error feedback."""
    if not isinstance(frame, pd.DataFrame):
        raise ExcelInputError("Excel 反馈数据必须为数据表。")
    source_rows = frame.attrs.get("source_rows", list(range(2, len(frame) + 2)))
    if len(source_rows) != len(frame) or any(type(row) is not int or row < 2 for row in source_rows) or len(set(source_rows)) != len(source_rows):
        raise ExcelInputError("Excel 原始行号无效，请重新读取工作表。")
    blocked = set(map(tuple, frame.attrs.get("formula_cells", []))) | set(map(tuple, frame.attrs.get("error_cells", [])))
    working = frame.copy(deep=True)
    if text_column in working.columns:
        column_position = working.columns.get_loc(text_column)
        for index, source_row in enumerate(source_rows):
            if (source_row, text_column) in blocked:
                # Numeric sentinel is skipped before duplicate detection. Its
                # original expression is restored only in the local row result.
                working.iat[index, column_position] = 0
    rows = prepare_feedback_rows(working, text_column, max_rows=max_rows)
    original_to_physical = {row.feedback_id: f"row-{physical}" for row, physical in zip(rows, source_rows)}
    for index, (row, source_row) in enumerate(zip(rows, source_rows)):
        row.feedback_id = f"row-{source_row}"
        row.source_row = source_row
        if row.duplicate_of:
            row.duplicate_of = original_to_physical[row.duplicate_of]
        if (source_row, text_column) in blocked:
            value = frame.iloc[index][text_column]
            row.original_text = str(value) if value is not None else ""
            row.processing_status = ProcessingStatus.SKIPPED_INVALID
            row.error_message = "评价单元格含公式或 Excel 错误，已跳过；请粘贴为文本值后重试。"
    return rows
