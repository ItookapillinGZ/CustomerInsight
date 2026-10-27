"""Deterministic, isolated demo reports and honest offline replay provenance."""
import io
from pathlib import Path
import tempfile
import tomllib
import unittest
from unittest.mock import patch
import zipfile

import pandas as pd
from openpyxl import load_workbook

from core.analysis_agent import AnalysisAgentResult, ReportAgentContext, report_fingerprint
from core.demo_workspace import _FIXTURE, load_demo_workspace, save_demo_workspace
from core.history_store import HistoryStore


class DemoWorkspaceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.workspace = load_demo_workspace()

    def test_reports_are_complete_two_batches_with_synthetic_provenance(self):
        self.assertEqual(self.workspace["data_origin"], "synthetic")
        self.assertEqual(self.workspace["metadata"]["model_calls"], 0)
        for key in ("current_report", "previous_report"):
            report = self.workspace[key]
            self.assertEqual(report["metrics"]["total_rows"], 30)
            self.assertEqual(report["metrics"]["success_rows"], 30)
            self.assertEqual(report["topic_coverage"]["status"], "complete")
            self.assertEqual(report["provenance"]["data_origin"], "synthetic")
            self.assertEqual(report["analysis_meta"]["label_origin"], "predefined")
            self.assertEqual(report["analysis_meta"]["model_calls"], 0)
            self.assertTrue(all("未调用模型" in source for source in report["detail"][report["detail_columns"]["result_source"]]))

    def test_fixture_lives_in_package_and_is_declared_as_wheel_data(self):
        root = Path(__file__).resolve().parents[1]
        self.assertEqual(_FIXTURE, root / "core" / "demo_data" / "reviews.json")
        self.assertTrue(_FIXTURE.is_file())
        config = tomllib.loads((root / "pyproject.toml").read_text(encoding="utf-8"))
        self.assertIn("demo_data/*.json", config["tool"]["setuptools"]["package-data"]["core"])

    def test_read_only_loader_does_not_touch_network_history_or_credentials(self):
        original_read = Path.read_text
        reads = []

        def read(path, *args, **kwargs):
            reads.append(str(path))
            return original_read(path, *args, **kwargs)

        with patch("pathlib.Path.read_text", read), \
             patch("socket.socket.connect", side_effect=AssertionError("No network")), \
             patch("sqlite3.connect", side_effect=AssertionError("No history")), \
             patch("core.analysis_agent.get_openai_client", side_effect=AssertionError("No model client")):
            workspace = load_demo_workspace()
        self.assertEqual(len(reads), 1)
        self.assertTrue(reads[0].endswith("reviews.json"))
        self.assertEqual(workspace["metadata"]["model_calls"], 0)

    def test_every_topic_is_grounded_in_its_own_batch_and_physical_row(self):
        for key in ("current_report", "previous_report"):
            report = self.workspace[key]
            names = report["detail_columns"]
            detail = report["detail"].set_index(names["source_row"])
            self.assertEqual(detail.index.tolist(), list(range(2, 32)))
            self.assertEqual(detail["评价编号"].nunique(), 30)
            self.assertEqual(detail[report["text_column"]].nunique(), 30)
            for record in report["aspects_by_row"]:
                for aspect in record["aspects"]:
                    self.assertIn(aspect["evidence"], detail.loc[record["source_row"], report["text_column"]])
            for topic in report["topics"]:
                self.assertEqual(topic["review_count"], len(set(topic["source_rows"])))
                for example in topic["examples"]:
                    self.assertIn(example["text"], detail.loc[example["source_row"], report["text_column"]])

    def test_comparison_is_recomputed_from_complete_same_schema_reports(self):
        current, previous = self.workspace["current_report"], self.workspace["previous_report"]
        comparison = current["historical_context"]["comparison"]
        self.assertTrue(comparison["topics_comparable"])
        self.assertTrue(comparison["basic_comparable"])
        metric = comparison["metrics"]["negative_rate"]
        self.assertAlmostEqual(metric["current"], current["metrics"]["negative_rows"] / 30)
        self.assertAlmostEqual(metric["previous"], previous["metrics"]["negative_rows"] / 30)
        self.assertEqual(current["period"]["start"], "2026-10-08")
        self.assertEqual(previous["period"]["end"], "2026-10-07")
        self.assertFalse(any("有重叠" in warning for warning in comparison["warnings"]))

    def test_replays_validate_and_never_claim_llm_generation_or_calls(self):
        fingerprint = report_fingerprint(self.workspace["current_report"])
        self.assertEqual(len(self.workspace["questions"]), 2)
        for question, payload in zip(self.workspace["questions"], self.workspace["replay_turns"]):
            turn = AnalysisAgentResult.model_validate(payload)
            self.assertEqual(turn.question, question)
            self.assertEqual(turn.report_fingerprint, fingerprint)
            self.assertEqual(turn.answer_origin, "fallback")
            self.assertEqual(turn.model_calls, 0)
            self.assertEqual(turn.tool_calls, 0)
            self.assertIn("离线示例回放", turn.scope_note)
            self.assertIn("未调用 LLM", turn.scope_note)

    def test_replay_evidence_separates_current_from_previous_even_for_same_row(self):
        current = self.workspace["current_report"]
        previous = self.workspace["previous_report"]
        names, old_names = current["detail_columns"], previous["detail_columns"]
        now_detail = current["detail"].set_index(names["source_row"])
        old_detail = previous["detail"].set_index(old_names["source_row"])
        turn = self.workspace["replay_turns"][1]
        self.assertTrue(turn["evidence_rows"])
        self.assertTrue(turn["historical_evidence"])
        for row in turn["evidence_rows"]:
            self.assertIn(row, now_detail.index)
        for evidence in turn["historical_evidence"]:
            self.assertEqual(evidence["batch_id"], self.workspace["previous_batch"]["id"])
            self.assertEqual(evidence["text"], old_detail.loc[evidence["source_row"], previous["text_column"]])
        self.assertIn("两个批次的行号独立", turn["conclusion"])

    def test_current_report_can_use_existing_table_and_history_agent_tools(self):
        report = self.workspace["current_report"]
        context = ReportAgentContext(report)
        overview = context.query("get_overview", {})
        self.assertIn("tool", overview)
        prior = context.query("read_previous_reviews", {"source_rows": [2], "limit": 1})
        self.assertTrue(prior["available"])
        self.assertEqual(prior["evidence"][0]["source_row"], 2)
        self.assertEqual(prior["batch"]["id"], self.workspace["previous_batch"]["id"])

    def test_workbook_has_raw_records_only_and_exact_row_identity(self):
        workbook = load_workbook(io.BytesIO(self.workspace["workbook_bytes"]), data_only=True)
        self.assertEqual(workbook.sheetnames, ["本批评价", "上批评价", "来源说明"])
        for sheet_name, report_key in (("本批评价", "current_report"), ("上批评价", "previous_report")):
            sheet = workbook[sheet_name]
            records = list(sheet.values)
            self.assertEqual(records[0], ("评价编号", "商品", "评价日期", "顾客评价"))
            self.assertEqual(len(records), 31)
            report = self.workspace[report_key]
            expected = report["detail"][["评价编号", "商品", "评价日期", "顾客评价"]]
            pd.testing.assert_frame_equal(pd.DataFrame(records[1:], columns=records[0]), expected)
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet.auto_filter.ref, "A1:D31")
        self.assertIn("固定合成示例", str(list(workbook["来源说明"].values)))
        workbook.close()

    def test_mapping_points_at_independent_sheets_in_shared_original(self):
        for key, sheet in (("current_mapping", "本批评价"), ("previous_mapping", "上批评价")):
            self.assertEqual(self.workspace[key], {"sheet_name": sheet, "text_column": "顾客评价",
                                                   "group_column": "商品", "date_column": "评价日期", "encoding": None})

    def test_original_workbook_bytes_are_stable_for_import_deduplication(self):
        again = load_demo_workspace()
        self.assertEqual(self.workspace["workbook_bytes"], again["workbook_bytes"])
        with zipfile.ZipFile(io.BytesIO(again["workbook_bytes"])) as archive:
            self.assertTrue(all(item.date_time == (2026, 10, 15, 0, 0, 0) for item in archive.infolist()))
            self.assertEqual(archive.read("docProps/core.xml").decode("utf-8").count("2026-10-15T00:00:00Z"), 2)

    def test_saved_demo_preserves_provenance_and_sheet_mappings_without_fixture_link(self):
        base = Path(__file__).resolve().parents[1] / ".test-tmp"
        base.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=base) as temporary:
            history = HistoryStore(Path(temporary) / "history")
            identifiers = save_demo_workspace(history)
            self.assertEqual(len(history.list_batches()), 2)
            self.assertNotEqual(identifiers["current_batch_id"], identifiers["previous_batch_id"])
            for kind in ("current", "previous"):
                record = history.load_batch(identifiers[f"{kind}_batch_id"])
                self.assertNotIn("historical_context", record["report"])
                self.assertEqual(record["report"]["provenance"]["data_origin"], "synthetic")
                self.assertEqual(record["report"]["analysis_meta"]["actual_requests"], 0)
                self.assertEqual(record["report"]["analysis_meta"]["demo_version"], "demo_workspace_v1")
                self.assertEqual(record["original_bytes"], self.workspace["workbook_bytes"])
                self.assertEqual(record["mapping"], self.workspace[f"{kind}_mapping"])

    def test_repeated_import_is_idempotent_and_does_not_overwrite_existing_chat(self):
        base = Path(__file__).resolve().parents[1] / ".test-tmp"
        base.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=base) as temporary:
            history = HistoryStore(Path(temporary) / "history")
            identifiers = save_demo_workspace(history)
            current_id = identifiers["current_batch_id"]
            record = history.load_batch(current_id)
            turn = dict(self.workspace["replay_turns"][0])
            turn["report_fingerprint"] = record["report_fingerprint"]
            chat = [{"result": turn}]
            history.save_chat(current_id, chat)
            with patch.object(history, "save_batch", side_effect=AssertionError("Existing imports must not be overwritten")):
                repeated = save_demo_workspace(history)
            self.assertEqual(identifiers, repeated)
            self.assertEqual(history.load_chat(current_id), chat)
            self.assertEqual(history.load_batch(current_id)["updated_at"], record["updated_at"])

    def test_import_uses_only_the_passed_history_and_does_not_share_ids(self):
        base = Path(__file__).resolve().parents[1] / ".test-tmp"
        base.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=base) as temporary:
            first = HistoryStore(Path(temporary) / "first")
            second = HistoryStore(Path(temporary) / "second")
            first_ids = save_demo_workspace(first)
            self.assertEqual(second.list_batches(), [])
            second_ids = save_demo_workspace(second)
            self.assertTrue(set(first_ids.values()).isdisjoint(second_ids.values()))
            self.assertEqual(len(first.list_batches()), 2)
            self.assertEqual(len(second.list_batches()), 2)

    def test_loading_twice_produces_fresh_equal_report_and_replay_objects(self):
        first, second = load_demo_workspace(), load_demo_workspace()
        self.assertEqual(report_fingerprint(first["current_report"]), report_fingerprint(second["current_report"]))
        self.assertEqual(first["replay_turns"], second["replay_turns"])
        first["current_report"]["detail"].iloc[0, 0] = "changed"
        first["replay_turns"][0]["conclusion"] = "changed"
        self.assertNotEqual(first["current_report"]["detail"].iloc[0, 0], second["current_report"]["detail"].iloc[0, 0])
        self.assertNotEqual(first["replay_turns"][0]["conclusion"], second["replay_turns"][0]["conclusion"])
        self.assertNotIn("historical_context", first["previous_report"])

    def test_replay_statistics_are_derived_from_real_fixture_counts(self):
        report = self.workspace["current_report"]
        top = max((item for item in report["topics"] if item["kind"] == "issue"), key=lambda item: (item["review_count"], item["topic"]))
        first, second = self.workspace["replay_turns"]
        self.assertIn(f"共 {top['review_count']} 条", first["conclusion"])
        self.assertIn(f"{top['share']:.1%}", first["conclusion"])
        self.assertIn(f"{top['review_count']}/30", second["conclusion_metrics"][1])
        self.assertEqual(report["metrics"]["negative_rows"], 15)
        self.assertEqual(self.workspace["previous_report"]["metrics"]["negative_rows"], 13)


if __name__ == "__main__":
    unittest.main()
